"""
utils/file_processors.py - File processing utilities
"""
from pathlib import Path
from typing import List, Dict, Tuple

from bs4 import BeautifulSoup


def process_file(file_path: Path, filename: str) -> Tuple[List[str], List[Dict]]:
    """Process a file and extract text chunks and metadata"""
    ext = file_path.suffix.lower()
    raw_text = ""
    
    try:
        if ext == ".pdf":
            raw_text = _process_pdf(file_path)
        elif ext in {".txt", ".md"}:
            raw_text = _process_text(file_path)
        elif ext == ".docx":
            raw_text = _process_docx(file_path)
        elif ext == ".pptx":
            raw_text = _process_pptx(file_path)
        elif ext == ".xlsx":
            raw_text = _process_xlsx(file_path)
        elif ext == ".csv":
            raw_text = _process_csv(file_path)
        elif ext in {".html", ".htm"}:
            raw_text = _process_html(file_path)
        else:
            return [], []
        
        # Import here to avoid circular dependency
        from ..vectorstore import chunk_text
        
        # Chunk the text
        chunks = chunk_text(raw_text)
        
        # Create metadata for each chunk
        metadatas = [{"source": filename} for _ in chunks]
        
        return chunks, metadatas
        
    except Exception as e:
        print(f"Error processing file {filename}: {str(e)}")
        return [], []


def _process_pdf(file_path: Path) -> str:
    """Process PDF file"""
    try:
        import pypdf
    except ImportError:
        raise ImportError("pypdf package not installed")
    
    reader = pypdf.PdfReader(str(file_path))
    text_parts = []
    
    for page in reader.pages:
        text = page.extract_text()
        if text:
            text_parts.append(text)
    
    return "\n".join(text_parts)


def _process_text(file_path: Path) -> str:
    """Process text file"""
    return file_path.read_text(errors='ignore')


def _process_docx(file_path: Path) -> str:
    """Process DOCX file"""
    try:
        import docx
    except ImportError:
        raise ImportError("python-docx package not installed")
    
    doc = docx.Document(str(file_path))
    return "\n".join(para.text for para in doc.paragraphs)


def _process_pptx(file_path: Path) -> str:
    """Process PPTX file"""
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError("python-pptx package not installed")
    
    prs = Presentation(str(file_path))
    text_parts = []
    
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text_parts.append(shape.text)
    
    return "\n".join(text_parts)


def _process_xlsx(file_path: Path) -> str:
    """Process XLSX file"""
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl package not installed")
    
    wb = openpyxl.load_workbook(str(file_path))
    text_parts = []
    
    for sheet in wb.worksheets:
        for row in sheet.rows:
            row_text = []
            for cell in row:
                if cell.value:
                    row_text.append(str(cell.value))
            if row_text:
                text_parts.append("\t".join(row_text))
    
    return "\n".join(text_parts)


def _process_csv(file_path: Path) -> str:
    """Process CSV file"""
    import csv
    
    text_parts = []
    with open(str(file_path), 'r', errors='ignore') as csvfile:
        reader = csv.reader(csvfile)
        for row in reader:
            text_parts.append(",".join(str(cell) for cell in row))
    
    return "\n".join(text_parts)


def _process_html(file_path: Path) -> str:
    """Process HTML file"""
    with open(str(file_path), 'r', errors='ignore') as htmlfile:
        soup = BeautifulSoup(htmlfile.read(), "html.parser")
        return soup.get_text("\n", strip=True)